from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from openai import OpenAI

from app.agents.resume_builder.contracts import ResumeSourceToMarkdownRequest, ResumeSourceToMarkdownResult
from app.agents.resume_builder.prompts import RESUME_SOURCE_TO_MARKDOWN_SYSTEM_PROMPT
from app.core.config import settings
from app.profile_service import extract_text_from_upload


class ResumeSourceToMarkdownAgentError(RuntimeError):
    """Raised when the resume-source-to-markdown agent cannot complete successfully."""


class ResumeSourceToMarkdownAgent:
    def __init__(
        self,
        *,
        client: OpenAI | Any | None = None,
        api_key: str | None = None,
        model: str | None = None,
    ) -> None:
        resolved_api_key = api_key or settings.openai_api_key
        self.model = model or settings.openai_resume_builder_resume_import_model
        if client is not None:
            self.client = client
            return
        if not resolved_api_key:
            raise ResumeSourceToMarkdownAgentError(
                "OPENAI_API_KEY is not configured. Set OPENAI_API_KEY before running the resume import agent."
            )
        self.client = OpenAI(api_key=resolved_api_key)

    def generate_markdown(self, request: ResumeSourceToMarkdownRequest) -> ResumeSourceToMarkdownResult:
        source_path = Path(request.input_resume_path)
        if not source_path.exists():
            raise ResumeSourceToMarkdownAgentError(f"Resume input file does not exist: {source_path}")
        if source_path.is_dir():
            raise ResumeSourceToMarkdownAgentError(f"Resume input path points to a directory: {source_path}")

        source_format = _detect_source_format(source_path)
        raw_text = _extract_resume_text(source_path, source_format)
        if not raw_text.strip():
            raise ResumeSourceToMarkdownAgentError("Could not extract any resume text from the source file.")

        response = self.client.responses.create(
            model=self.model,
            input=[
                {
                    "role": "system",
                    "content": [{"type": "input_text", "text": RESUME_SOURCE_TO_MARKDOWN_SYSTEM_PROMPT}],
                },
                {
                    "role": "user",
                    "content": [{"type": "input_text", "text": _render_user_prompt(raw_text, source_format)}],
                },
            ],
        )
        markdown = _extract_output_text(response).strip()
        if not markdown:
            raise ResumeSourceToMarkdownAgentError("OpenAI response did not include normalized resume markdown.")

        return ResumeSourceToMarkdownResult(
            input_resume_path=str(source_path),
            output_markdown_path=request.output_markdown_path,
            source_format=source_format,
            markdown_content=markdown,
        )


def _detect_source_format(source_path: Path) -> str:
    suffix = source_path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    raise ResumeSourceToMarkdownAgentError(
        f"Unsupported resume file type: {source_path.suffix or '<none>'}. Only .pdf and .docx are supported."
    )


def _extract_resume_text(source_path: Path, source_format: str) -> str:
    if source_format == "pdf":
        return extract_text_from_upload(source_path.name, source_path.read_bytes())
    if source_format == "docx":
        return _extract_text_from_docx(source_path)
    raise ResumeSourceToMarkdownAgentError(f"Unsupported resume source format: {source_format}")


def _extract_text_from_docx(source_path: Path) -> str:
    document = Document(str(source_path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(text)
    return "\n".join(lines).strip()


def _render_user_prompt(raw_text: str, source_format: str) -> str:
    return (
        "Convert the following extracted resume text into clean markdown.\n\n"
        f"Source format: {source_format}\n\n"
        "Requirements:\n"
        "- Preserve the candidate's actual content.\n"
        "- Remove obvious extraction artifacts.\n"
        "- Organize the resume into clear markdown sections.\n"
        "- Do not invent content.\n\n"
        "## Extracted Resume Text\n"
        f"{raw_text}\n"
    )


def _extract_output_text(response: Any) -> str:
    output_text = getattr(response, "output_text", None)
    if isinstance(output_text, str) and output_text.strip():
        return output_text

    output = getattr(response, "output", None)
    if not isinstance(output, list):
        return ""

    collected: list[str] = []
    for item in output:
        content_items = getattr(item, "content", None)
        if not isinstance(content_items, list):
            continue
        for content in content_items:
            content_type = getattr(content, "type", None)
            text = getattr(content, "text", None)
            if content_type == "output_text" and isinstance(text, str) and text.strip():
                collected.append(text)
    return "\n\n".join(collected)
