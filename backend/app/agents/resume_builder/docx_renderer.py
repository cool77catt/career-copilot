from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph

from app.agents.resume_builder.contracts import ResumeDocxRenderRequest, ResumeDocxRenderResult


class ResumeDocxRendererError(RuntimeError):
    """Raised when the markdown-to-docx renderer cannot complete successfully."""


@dataclass(frozen=True)
class _MarkdownItem:
    kind: str
    text: str


@dataclass(frozen=True)
class _MarkdownSection:
    title: str
    items: list[_MarkdownItem]


@dataclass(frozen=True)
class _TemplateSection:
    heading: Paragraph
    body_paragraphs: list[Paragraph]


class ResumeDocxRenderer:
    def render(self, request: ResumeDocxRenderRequest) -> ResumeDocxRenderResult:
        template_path = Path(request.template_docx_path)
        if not template_path.exists():
            raise ResumeDocxRendererError(f"Template .docx file does not exist: {template_path}")
        if template_path.is_dir():
            raise ResumeDocxRendererError(f"Template .docx path points to a directory: {template_path}")

        output_path = Path(request.output_docx_path)
        if output_path.exists() and output_path.is_dir():
            raise ResumeDocxRendererError(f"Output .docx path points to a directory: {output_path}")

        markdown_sections = _parse_markdown_sections(request.resume_markdown)
        if not markdown_sections:
            raise ResumeDocxRendererError("Resume markdown did not include any renderable sections.")

        document = Document(str(template_path))
        template_sections = _extract_template_sections(document, markdown_sections)
        if not template_sections:
            raise ResumeDocxRendererError("Template .docx does not include any matching section headings to populate.")

        notes = _apply_sections(document, template_sections, markdown_sections)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))

        return ResumeDocxRenderResult(
            output_docx_path=str(output_path),
            template_docx_path=str(template_path),
            formatting_notes=notes,
        )


def _parse_markdown_sections(markdown: str) -> list[_MarkdownSection]:
    sections: list[_MarkdownSection] = []
    current_title: str | None = None
    current_items: list[_MarkdownItem] = []

    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            if current_title is not None:
                sections.append(_MarkdownSection(title=_strip_markdown(current_title), items=current_items))
            current_title = line[3:].strip()
            current_items = []
            continue
        if current_title is None:
            continue
        current_items.append(_parse_markdown_item(line))

    if current_title is not None:
        sections.append(_MarkdownSection(title=_strip_markdown(current_title), items=current_items))

    return sections


def _parse_markdown_item(line: str) -> _MarkdownItem:
    if line.startswith("### "):
        return _MarkdownItem(kind="subheading", text=_strip_markdown(line[4:].strip()))
    if line.startswith("- "):
        return _MarkdownItem(kind="bullet", text=_strip_markdown(line[2:].strip()))
    return _MarkdownItem(kind="paragraph", text=_strip_markdown(line))


def _extract_template_sections(
    document: DocumentType,
    markdown_sections: list[_MarkdownSection],
) -> list[_TemplateSection]:
    paragraphs = list(document.paragraphs)
    markdown_titles = {_normalize_heading(section.title) for section in markdown_sections}
    heading_indexes = [
        index for index, paragraph in enumerate(paragraphs) if _normalize_heading(paragraph.text) in markdown_titles
    ]
    sections: list[_TemplateSection] = []
    for offset, index in enumerate(heading_indexes):
        next_index = heading_indexes[offset + 1] if offset + 1 < len(heading_indexes) else len(paragraphs)
        body = [paragraphs[position] for position in range(index + 1, next_index)]
        sections.append(_TemplateSection(heading=paragraphs[index], body_paragraphs=body))
    return sections


def _apply_sections(
    document: DocumentType,
    template_sections: list[_TemplateSection],
    markdown_sections: list[_MarkdownSection],
) -> list[str]:
    notes: list[str] = []
    mapping = _build_section_mapping(template_sections, markdown_sections)
    used_markdown_indexes = set(mapping.values())

    for template_index, template_section in enumerate(template_sections):
        markdown_index = mapping.get(template_index)
        if markdown_index is None:
            notes.append(f"Preserved unmatched template section: {template_section.heading.text.strip()}")
            continue

        markdown_section = markdown_sections[markdown_index]
        paragraph_template, bullet_template = _select_body_templates(template_section.body_paragraphs)
        insertion_point = _next_heading_paragraph(template_sections, template_index)
        existing_body = list(template_section.body_paragraphs)

        for item_index, item in enumerate(markdown_section.items):
            chosen_template = bullet_template if item.kind == "bullet" else paragraph_template
            if chosen_template is None:
                chosen_template = paragraph_template or bullet_template

            if item_index < len(existing_body):
                paragraph = existing_body[item_index]
            elif chosen_template is not None:
                paragraph = _clone_paragraph_before(document, chosen_template, insertion_point)
            else:
                paragraph = _insert_plain_paragraph(document, insertion_point)

            _replace_paragraph_text(paragraph, item.text, force_bold=item.kind == "subheading")

        for paragraph in existing_body[len(markdown_section.items) :]:
            _remove_paragraph(paragraph)

    remaining_sections = [
        markdown_sections[index] for index in range(len(markdown_sections)) if index not in used_markdown_indexes
    ]
    if remaining_sections:
        notes.append("Skipped markdown sections without matching headings in the template to preserve layout fidelity.")

    notes.append("Preserved template header and section headings; rewrote section bodies using template paragraph clones.")
    return notes


def _build_section_mapping(
    template_sections: list[_TemplateSection],
    markdown_sections: list[_MarkdownSection],
) -> dict[int, int]:
    mapping: dict[int, int] = {}
    available_markdown = set(range(len(markdown_sections)))

    for template_index, template_section in enumerate(template_sections):
        template_title = _normalize_heading(template_section.heading.text)
        for markdown_index in sorted(available_markdown):
            if _normalize_heading(markdown_sections[markdown_index].title) == template_title:
                mapping[template_index] = markdown_index
                available_markdown.remove(markdown_index)
                break

    return mapping


def _normalize_heading(text: str) -> str:
    return " ".join(text.lower().split())


def _select_body_templates(body_paragraphs: list[Paragraph]) -> tuple[Paragraph | None, Paragraph | None]:
    paragraph_template: Paragraph | None = None
    bullet_template: Paragraph | None = None
    for paragraph in body_paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if "Bullet" in style_name and bullet_template is None:
            bullet_template = paragraph
        elif paragraph.text.strip() and paragraph_template is None:
            paragraph_template = paragraph
    return paragraph_template, bullet_template


def _next_heading_paragraph(template_sections: list[_TemplateSection], current_index: int) -> Paragraph | None:
    next_index = current_index + 1
    if next_index >= len(template_sections):
        return None
    return template_sections[next_index].heading


def _clone_paragraph_before(
    document: DocumentType,
    template_paragraph: Paragraph,
    insertion_point: Paragraph | None,
) -> Paragraph:
    cloned_element = deepcopy(template_paragraph._element)
    _clear_paragraph_element_runs(cloned_element)
    if insertion_point is None:
        document._body._element.append(cloned_element)
    else:
        insertion_point._element.addprevious(cloned_element)
    return Paragraph(cloned_element, template_paragraph._parent)


def _insert_plain_paragraph(document: DocumentType, insertion_point: Paragraph | None) -> Paragraph:
    if insertion_point is None:
        paragraph = document.add_paragraph()
    else:
        paragraph = insertion_point.insert_paragraph_before()
    paragraph.style = "Normal"
    return paragraph


def _replace_paragraph_text(paragraph: Paragraph, text: str, *, force_bold: bool = False) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    _write_markdown_runs(paragraph, text, force_bold=force_bold)


def _write_markdown_runs(paragraph: Paragraph, text: str, *, force_bold: bool = False) -> None:
    segments = _split_markdown_segments(text)
    if not segments:
        paragraph.add_run("")
        return
    for segment_text, is_bold in segments:
        run = paragraph.add_run(segment_text)
        if force_bold or is_bold:
            run.bold = True


def _split_markdown_segments(text: str) -> list[tuple[str, bool]]:
    clean_text = text.replace("\t", " ").strip()
    if not clean_text:
        return []
    parts = re.split(r"(\*\*.*?\*\*)", clean_text)
    segments: list[tuple[str, bool]] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            segments.append((part[2:-2], True))
        else:
            segments.append((part, False))
    return segments


def _strip_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "")
    text = text.replace("*", "")
    return " ".join(text.split()).strip()


def _clear_paragraph_element_runs(paragraph_element: CT_P) -> None:
    for child in list(paragraph_element):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            paragraph_element.remove(child)


def _remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
