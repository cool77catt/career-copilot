import json
from types import SimpleNamespace

import pytest
from docx import Document

from app.agents.resume_builder.contracts import ResumeSourceToMarkdownRequest
from app.agents.resume_builder.resume_source_to_markdown_agent import (
    ResumeSourceToMarkdownAgent,
    ResumeSourceToMarkdownAgentError,
    _extract_text_from_docx,
)
from app.agents.resume_builder.resume_source_to_markdown_cli import main
from app.agents.resume_builder.resume_source_to_markdown_runner import run_resume_source_to_markdown
from app.core.config import settings


class _FakeResponsesClient:
    def __init__(self, output_text: str) -> None:
        self.output_text = output_text
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(output_text=self.output_text)


class _FakeOpenAIClient:
    def __init__(self, output_text: str) -> None:
        self.responses = _FakeResponsesClient(output_text)


def test_docx_text_extraction_reads_paragraphs_and_tables(tmp_path):
    source_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Staff Software Engineer")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "jane@example.com"
    table.cell(0, 1).text = "linkedin.com/in/janedoe"
    document.save(str(source_path))

    extracted = _extract_text_from_docx(source_path)

    assert "Jane Doe" in extracted
    assert "jane@example.com" in extracted
    assert "linkedin.com/in/janedoe" in extracted


def test_resume_source_agent_converts_docx_to_markdown(tmp_path):
    source_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Built backend systems.")
    document.save(str(source_path))

    client = _FakeOpenAIClient("# Resume\n\n## Summary\nBuilt backend systems.")
    agent = ResumeSourceToMarkdownAgent(client=client, model="test-model")

    result = agent.generate_markdown(
        ResumeSourceToMarkdownRequest(
            input_resume_path=str(source_path),
            output_markdown_path=str(tmp_path / "resume.md"),
        )
    )

    assert result.source_format == "docx"
    assert result.markdown_content.startswith("# Resume")
    assert client.responses.calls[0]["model"] == "test-model"


def test_resume_source_agent_requires_api_key_without_injected_client(monkeypatch):
    monkeypatch.setattr(settings, "openai_api_key", None, raising=False)

    with pytest.raises(ResumeSourceToMarkdownAgentError):
        ResumeSourceToMarkdownAgent()


def test_resume_source_agent_rejects_unsupported_extension(tmp_path):
    source_path = tmp_path / "resume.txt"
    source_path.write_text("Resume text", encoding="utf-8")
    client = _FakeOpenAIClient("# Resume")
    agent = ResumeSourceToMarkdownAgent(client=client)

    with pytest.raises(ResumeSourceToMarkdownAgentError):
        agent.generate_markdown(
            ResumeSourceToMarkdownRequest(
                input_resume_path=str(source_path),
                output_markdown_path=str(tmp_path / "resume.md"),
            )
        )


def test_resume_source_runner_writes_markdown_file(tmp_path):
    source_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.save(str(source_path))

    client = _FakeOpenAIClient("# Resume\n\n## Summary\nImported resume.")
    agent = ResumeSourceToMarkdownAgent(client=client)
    output_path = tmp_path / "resume.md"

    result = run_resume_source_to_markdown(
        ResumeSourceToMarkdownRequest(
            input_resume_path=str(source_path),
            output_markdown_path=str(output_path),
        ),
        agent=agent,
    )

    assert result.output_markdown_path == str(output_path)
    assert output_path.read_text(encoding="utf-8").startswith("# Resume")


def test_resume_source_cli_with_mocked_runner_prints_json(tmp_path, monkeypatch, capsys):
    input_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Jane Doe")
    document.save(str(input_path))
    output_path = tmp_path / "resume.md"

    monkeypatch.setattr(
        "app.agents.resume_builder.resume_source_to_markdown_cli.run_resume_source_to_markdown",
        lambda request: SimpleNamespace(
            model_dump=lambda: {
                "input_resume_path": request.input_resume_path,
                "output_markdown_path": request.output_markdown_path,
                "source_format": "docx",
                "markdown_content": "# Resume\n",
            }
        ),
    )

    exit_code = main(
        [
            "--input-resume-file",
            str(input_path),
            "--output-markdown-path",
            str(output_path),
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out.strip())
    assert payload["source_format"] == "docx"
    assert payload["output_markdown_path"] == str(output_path)
