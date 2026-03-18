import json
from pathlib import Path

import pytest
from docx import Document

from app.agents.resume_builder.contracts import ResumeDocxRenderRequest
from app.agents.resume_builder.docx_renderer import ResumeDocxRenderer, ResumeDocxRendererError
from app.agents.resume_builder.docx_runner import run_resume_docx_renderer
from app.agents.resume_builder.render_docx_cli import main


def _build_template_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Jane Doe", style="Heading 1")
    document.add_paragraph("Staff Backend Engineer")
    document.add_paragraph("jane@example.com")
    document.add_paragraph("Summary", style="Heading 1")
    document.add_paragraph("Old summary", style="Normal")
    document.add_paragraph("Experience", style="Heading 1")
    document.add_paragraph("Old experience bullet", style="List Bullet")
    document.add_paragraph("Skills", style="Heading 1")
    document.add_paragraph("Old skills", style="Normal")
    document.save(str(path))


def test_docx_renderer_generates_new_docx_with_template_structure(tmp_path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "output.docx"
    _build_template_docx(template_path)

    renderer = ResumeDocxRenderer()
    result = renderer.render(
        ResumeDocxRenderRequest(
            resume_markdown=(
                "# Jane Doe\n\n"
                "Staff Backend Engineer\n\n"
                "## Summary\n"
                "Backend engineer with **platform delivery** experience.\n\n"
                "## Experience\n"
                "### Technical Manager | Staff Software Engineer\n"
                "**Invisible AI** | Sept 2023 - Present\n"
                "- Built FastAPI APIs\n"
                "- Improved service reliability\n\n"
                "## Skills\n"
                "Python\n"
                "AWS\n"
            ),
            template_docx_path=str(template_path),
            output_docx_path=str(output_path),
        )
    )

    assert result.output_docx_path == str(output_path)
    assert output_path.exists()
    assert template_path.exists()

    rendered = Document(str(output_path))
    paragraphs = [paragraph.text for paragraph in rendered.paragraphs if paragraph.text.strip()]
    assert "Jane Doe" in paragraphs
    assert "Staff Backend Engineer" in paragraphs
    assert "jane@example.com" in paragraphs
    assert "Summary" in paragraphs
    assert "Backend engineer with platform delivery experience." in paragraphs
    assert "Technical Manager | Staff Software Engineer" in paragraphs
    assert "Invisible AI | Sept 2023 - Present" in paragraphs
    assert "Built FastAPI APIs" in paragraphs
    assert "Improved service reliability" in paragraphs
    assert "Python" in paragraphs
    assert "AWS" in paragraphs
    assert not any("**" in paragraph for paragraph in paragraphs)
    assert not any("###" in paragraph for paragraph in paragraphs)

    template = Document(str(template_path))
    template_paragraphs = [paragraph.text for paragraph in template.paragraphs if paragraph.text.strip()]
    assert "Old summary" in template_paragraphs
    assert "Old experience bullet" in template_paragraphs


def test_docx_runner_rejects_missing_template(tmp_path):
    with pytest.raises(ResumeDocxRendererError):
        run_resume_docx_renderer(
            ResumeDocxRenderRequest(
                resume_markdown="# Resume\n\n## Summary\nTest\n",
                template_docx_path=str(tmp_path / "missing.docx"),
                output_docx_path=str(tmp_path / "output.docx"),
            )
        )


def test_docx_cli_reads_files_and_prints_json(tmp_path, capsys):
    template_path = tmp_path / "template.docx"
    markdown_path = tmp_path / "resume.md"
    output_path = tmp_path / "output.docx"
    _build_template_docx(template_path)
    markdown_path.write_text(
        "# Resume\n\n## Summary\nPlatform engineer.\n\n## Experience\n- Built services\n",
        encoding="utf-8",
    )

    exit_code = main(
        [
            "--resume-markdown-file",
            str(markdown_path),
            "--template-docx-file",
            str(template_path),
            "--output-docx-path",
            str(output_path),
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out.strip())
    assert payload["output_docx_path"] == str(output_path)
    assert output_path.exists()


def test_docx_cli_fails_when_template_missing(tmp_path, capsys):
    markdown_path = tmp_path / "resume.md"
    output_path = tmp_path / "output.docx"
    markdown_path.write_text("# Resume\n\n## Summary\nPlatform engineer.\n", encoding="utf-8")

    exit_code = main(
        [
            "--resume-markdown-file",
            str(markdown_path),
            "--template-docx-file",
            str(tmp_path / "missing.docx"),
            "--output-docx-path",
            str(output_path),
        ]
    )

    assert exit_code == 1
    assert "Template .docx file does not exist" in capsys.readouterr().err
